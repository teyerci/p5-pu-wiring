"""Markdown worksheet -> one-page landscape docx (compact table)."""
import re, sys
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src, out = sys.argv[1], sys.argv[2]
FS = float(sys.argv[3]) if len(sys.argv) > 3 else 7
lines = open(src).read().splitlines()
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
for m in ("left_margin", "right_margin"):
    setattr(sec, m, Cm(1.0))
sec.top_margin = sec.bottom_margin = Cm(0.9)
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(FS)
st.paragraph_format.space_after = Pt(1); st.paragraph_format.space_before = Pt(0); st.paragraph_format.line_spacing = 1.0

def add_runs(par, text, size=None, bold_all=False):
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        b = part.startswith("**")
        t = part[2:-2] if b else part[1:-1] if part.startswith("`") else part
        r = par.add_run(t); r.bold = b or bold_all
        if size: r.font.size = Pt(size)

rows = [l for l in lines if l.startswith("|")]
header = [c.strip() for c in rows[0].strip("|").split("|")]
body = [[c.strip() for c in r.strip("|").split("|")] for r in rows[2:]]
title = lines[0].lstrip("# ").strip()
p = doc.add_paragraph(); add_runs(p, title, size=FS + 4, bold_all=True)
intro = next(l for l in lines[1:] if l.strip() and not l.startswith("|"))
p = doc.add_paragraph(); add_runs(p, intro)
widths = [Cm(1.0), Cm(7.4), Cm(1.7), Cm(7.5), Cm(10.1)]
tbl = doc.add_table(rows=1, cols=len(header)); tbl.style = "Table Grid"; tbl.autofit = False
for i, h in enumerate(header):
    c = tbl.rows[0].cells[i]; c.width = widths[i]; c.text = ""
    add_runs(c.paragraphs[0], h, bold_all=True)
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "E8EEF4"); c._tc.get_or_add_tcPr().append(shd)
for r in body:
    cells = tbl.add_row().cells
    for i, v in enumerate(r):
        cells[i].width = widths[i]; cells[i].text = ""
        add_runs(cells[i].paragraphs[0], v)
for i, w in enumerate(widths):
    tbl.columns[i].width = w
grid = tbl._tbl.tblGrid
for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
    gc.set(qn("w:w"), str(int(widths[i].twips)))
tblPr = tbl._tbl.tblPr
lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
after = lines[lines.index(rows[-1]) + 1:]
for l in after:
    if l.strip():
        p = doc.add_paragraph(); add_runs(p, l.lstrip("- "))
doc.save(out)
