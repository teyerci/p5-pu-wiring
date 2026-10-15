"""One-page summary md (intro, one 3-column table, closing lines) -> portrait docx."""
import re, sys
from docx import Document
from docx.shared import Cm, Pt
src, out, FS = sys.argv[1], sys.argv[2], float(sys.argv[3])
lines = open(src).read().splitlines()
doc = Document(); sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for m in ("left_margin","right_margin","top_margin","bottom_margin"): setattr(sec, m, Cm(1.6))
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(FS)
st.paragraph_format.space_after = Pt(4); st.paragraph_format.line_spacing = 1.05
def runs(p, t, size=None, bold=False):
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", t):
        if not part: continue
        b = part.startswith("**"); x = part[2:-2] if b else part[1:-1] if part.startswith("`") else part
        r = p.add_run(x); r.bold = b or bold
        if size: r.font.size = Pt(size)
rows = [l for l in lines if l.startswith("|")]
hdr = [c.strip() for c in rows[0].strip("|").split("|")]
body = [[c.strip() for c in r.strip("|").split("|")] for r in rows[2:]]
p = doc.add_paragraph(); runs(p, lines[0].lstrip("# "), FS+4, True)
done_table = False
for l in lines[1:]:
    if l.startswith("|"):
        if done_table: continue
        done_table = True
        t = doc.add_table(rows=1, cols=3); t.style = "Table Grid"; t.autofit = False
        ws = [Cm(3.0), Cm(6.6), Cm(8.2)]
        for i,h in enumerate(hdr):
            c = t.rows[0].cells[i]; c.width = ws[i]; runs(c.paragraphs[0], h, bold=True)
        for r in body:
            cells = t.add_row().cells
            for i,v in enumerate(r):
                cells[i].width = ws[i]; runs(cells[i].paragraphs[0], v, bold=(i==0))
        doc.add_paragraph()
    elif l.strip():
        p = doc.add_paragraph(); runs(p, l)
doc.save(out)
